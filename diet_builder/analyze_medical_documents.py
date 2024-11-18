import os
from pathlib import Path
import pytesseract
from PIL import Image
import PyPDF2
import docx
import openpyxl

def extract_text(file_path):
    """
    Extract text from a document.
    
    Supported formats: .txt, .pdf, .docx, .xlsx, .png, .jpg, .jpeg
    
    :param file_path: Path to the file.
    :return: Extracted text as a string.
    """
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")
    
    ext = file_path.suffix.lower()
    
    if ext == ".txt":
        # Handle plain text files
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    
    elif ext == ".pdf":
        # Handle PDF files
        text = ""
        with open(file_path, "rb") as f:
            pdf_reader = PyPDF2.PdfReader(f)
            for page in pdf_reader.pages:
                text += page.extract_text() or ""
        return text
    
    elif ext == ".docx":
        # Handle Word documents
        doc = docx.Document(file_path)
        return "\n".join([paragraph.text for paragraph in doc.paragraphs])
    
    elif ext == ".xlsx":
        # Handle Excel files
        wb = openpyxl.load_workbook(file_path)
        text = ""
        for sheet in wb:
            for row in sheet.iter_rows(values_only=True):
                text += " ".join([str(cell) for cell in row if cell]) + "\n"
        return text.strip()
    
    elif ext in [".png", ".jpg", ".jpeg"]:
        # Handle image files (OCR)
        image = Image.open(file_path)
        return pytesseract.image_to_string(image)
    
    else:
        raise ValueError(f"Unsupported file format: {ext}")

if __name__ == '__main__':
    # Example Usage
    file_path = "example.docx"
    print(extract_text(file_path))
